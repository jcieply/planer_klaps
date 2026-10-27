import re, textwrap, sys
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

src = open('/home/claude/klaps/src/sample.txt', encoding='utf-8').read().strip('\n')
blocks = [b.strip('\n') for b in re.split(r'\n\s*\n', src) if b.strip()]

# -------- klasyfikacja prostych bloków --------
items = []  # (type, lines)
for b in blocks:
    lines = b.split('\n')
    first = lines[0]
    if re.match(r'^\d+\.\s+(WN|PL)', first):
        items.append(('heading', [first]))
    elif first.strip() in ('KONIEC',):
        items.append(('end', lines))
    elif re.match(r'^[A-ZĄĆĘŁŃÓŚŹŻ ]+:$', first.strip()):
        items.append(('trans', lines))
    elif len(lines) > 1 and re.match(r'^[A-ZĄĆĘŁŃÓŚŹŻ ]+(\s*\([^)]*\))?$', first.strip()) and len(first) < 40:
        items.append(('dialog', lines))
    else:
        items.append(('action', lines))

# -------- PDF --------
pdfmetrics.registerFont(TTFont('Mono', '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'))
W, H = A4
LEFT, TOP, BOTTOM, LH, FS = 108, 72, 72, 12, 12
CW = FS * 0.602
c = canvas.Canvas('/home/claude/klaps/test/sample.pdf', pagesize=A4)
page = 1
y = H - TOP
def newpage():
    global page, y
    c.showPage(); page += 1; y = H - TOP
    c.setFont('Mono', FS)
    c.drawRightString(W - 72, H - 36, f'{page}.')
def need(n):
    global y
    if y - n * LH < BOTTOM: newpage()
c.setFont('Mono', FS)
# strona tytułowa
c.drawCentredString(W/2, H/2 + 40, items[0][1][0])
c.drawCentredString(W/2, H/2 + 10, items[1][1][0] if len(items) > 1 and items[1][0]=='action' else '')
newpage()
body = items[2:] if items[1][0]=='action' else items[1:]
for t, lines in body:
    if t == 'heading':
        need(3)
        c.drawString(LEFT, y, lines[0]); y -= 2 * LH
    elif t == 'action':
        text = ' '.join(lines)
        wrapped = textwrap.wrap(text, 60)
        need(min(len(wrapped), 3))
        for w in wrapped:
            need(1)
            c.drawString(LEFT, y, w); y -= LH
        y -= LH
    elif t == 'dialog':
        need(3)
        cue = lines[0]
        c.drawString(LEFT + 22 * CW, y, cue); y -= LH
        for ln in lines[1:]:
            if ln.startswith('('):
                need(1); c.drawString(LEFT + 16 * CW, y, ln); y -= LH
            else:
                for w in textwrap.wrap(ln, 35):
                    need(1); c.drawString(LEFT + 10 * CW, y, w); y -= LH
        y -= LH
    elif t == 'trans':
        need(2)
        c.drawRightString(W - 72, y, lines[0]); y -= 2 * LH
    elif t == 'end':
        need(2)
        c.drawCentredString(W/2, y, lines[0]); y -= 2 * LH
c.save()
print('pdf pages', page)

# -------- FDX --------
def esc(x): return x.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
paras = []
n = 0
for t, lines in body:
    if t == 'heading':
        n += 1
        txt = re.sub(r'^\d+\.\s+', '', lines[0])
        paras.append(f'<Paragraph Number="{n}" Type="Scene Heading"><Text>{esc(txt)}</Text></Paragraph>')
    elif t == 'action':
        paras.append(f'<Paragraph Type="Action"><Text>{esc(" ".join(lines))}</Text></Paragraph>')
    elif t == 'dialog':
        paras.append(f'<Paragraph Type="Character"><Text>{esc(lines[0])}</Text></Paragraph>')
        for ln in lines[1:]:
            ty = 'Parenthetical' if ln.startswith('(') else 'Dialogue'
            paras.append(f'<Paragraph Type="{ty}"><Text>{esc(ln)}</Text></Paragraph>')
    elif t == 'trans':
        paras.append(f'<Paragraph Type="Transition"><Text>{esc(lines[0])}</Text></Paragraph>')
    elif t == 'end':
        paras.append(f'<Paragraph Type="Action"><Text>{esc(lines[0])}</Text></Paragraph>')
fdx = '<?xml version="1.0" encoding="UTF-8" standalone="no" ?>\n<FinalDraft DocumentType="Script" Template="No" Version="2"><Content>\n' + '\n'.join(paras) + '\n</Content></FinalDraft>'
open('/home/claude/klaps/test/sample.fdx', 'w', encoding='utf-8').write(fdx)

# -------- DOCX (każdy akapit osobno, bez pustych linii) --------
from docx import Document
d = Document()
d.add_paragraph(items[0][1][0])
for t, lines in body:
    if t == 'heading':
        d.add_paragraph(lines[0])
    elif t == 'action':
        d.add_paragraph(' '.join(lines))
    elif t == 'dialog':
        for ln in lines: d.add_paragraph(ln)
    else:
        d.add_paragraph(lines[0])
d.save('/home/claude/klaps/test/sample.docx')
print('ok', len(body))
